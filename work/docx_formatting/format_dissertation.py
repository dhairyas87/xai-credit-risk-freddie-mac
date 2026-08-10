from copy import deepcopy
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.shared import RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = "work/docx_formatting/2024AB05218_Formatted_source.docx"
OUTPUT = "work/docx_formatting/2024AB05218_Formatted_final.docx"

TABLE_CAPTIONS = [
    ("Table 3.1", "Research Tools and Development Environment"),
    ("Table 4.1", "Feature Categories Used in the Proposed Framework"),
    ("Table 5.1", "Experimental Configuration"),
    ("Table 5.2", "Loan Amount Recommendation Performance"),
    ("Table 5.3", "Interest Rate Recommendation Performance"),
    ("Table 5.4", "Loan Term Recommendation Performance"),
    ("Table 5.5", "Average Borrower Characteristics by Stress Status"),
    ("Table 5.6", "Overall Model Comparison"),
    ("Table 5.7", "Summary of Key Findings"),
    ("Table A.1", "Borrower Features"),
    ("Table A.2", "Property Features"),
    ("Table A.3", "Loan Features"),
    ("Table B.1", "Models Evaluated"),
    ("Table B.2", "Stacked Ensemble Configuration"),
    ("Table C.1", "Borrower Profile"),
    ("Table C.2", "Property Information"),
    ("Table C.3", "Framework Output"),
    ("Table G.1", "Glossary of Terms"),
]

FIGURE_CAPTIONS = [
    ("Figure 3.1", "Overall Research Methodology"),
    ("Figure 3.2", "Prepared Dataset Directory Structure"),
    ("Figure 3.3", "Overall Research Pipeline"),
    ("Figure 4.1", "Overall Proposed Explainable Mortgage Lending Framework"),
    ("Figure 4.2", "Feature Categories within the Proposed Framework"),
    ("Figure 4.3", "Feature Store Directory Structure"),
    ("Figure 5.6", "Local SHAP Explanation"),
    ("Figure 5.7", "Partial Dependence Plot"),
    ("Figure 5.9", "LoanFit Recommendation Dashboard"),
    ("Figure 5.10", "LoanFit Explainability Dashboard"),
]


def iter_block_items(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def clear_paragraph_text(paragraph):
    for run in paragraph.runs:
        run.text = ""


def set_paragraph_text(paragraph, text):
    clear_paragraph_text(paragraph)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def format_caption(paragraph, number, title):
    set_paragraph_text(paragraph, f"{number}: {title}")
    if "Caption" in [style.name for style in paragraph.part.document.styles]:
        paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        run.italic = True
        run.bold = False
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)


def ensure_caption_style(doc):
    if "Caption" in [style.name for style in doc.styles]:
        return
    style = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    style.font.italic = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(6)


def paragraph_after(doc, paragraph, text="", style=None):
    new_para = doc.add_paragraph(text, style=style)
    paragraph._p.addnext(new_para._p)
    return new_para


def paragraph_before_table(doc, table, text="", style=None):
    new_para = doc.add_paragraph(text, style=style)
    table._tbl.addprevious(new_para._p)
    return new_para


def paragraph_before_paragraph(doc, paragraph, text="", style=None):
    new_para = doc.add_paragraph(text, style=style)
    paragraph._p.addprevious(new_para._p)
    return new_para


def has_drawing(paragraph):
    return bool(paragraph._p.xpath(".//w:drawing") or paragraph._p.xpath(".//w:pict"))


def normalize_styles(doc):
    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text and paragraph.style.name.startswith("Heading"):
            paragraph.style = "Normal"
            continue
        if paragraph.style.name == "Heading 1" and re.match(r"^\d+\.\d+\s+", text):
            paragraph.style = "Heading 2"
        if paragraph.style.name == "Heading 1" and re.match(r"^\d+\.\d+\.\d+\s+", text):
            paragraph.style = "Heading 3"
        if re.match(r"^\d+\.\d+\.\d+\s+", text):
            paragraph.style = "Heading 3"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.italic = False
                run.font.color.rgb = RGBColor(21, 86, 110)
        if text == "Discussion":
            paragraph.style = "Heading 3"


def set_major_section_page_breaks(doc):
    section_starts = {
        "Acknowledgement",
        "Abstract",
        "Chapter 1",
        "CHAPTER 2",
        "CHAPTER 3",
        "CHAPTER 4",
        "CHAPTER 5",
        "CHAPTER 6",
        "APPENDICES",
        "Appendix B: Machine Learning Model Configuration",
        "Appendix C: Sample Recommendation",
        "GLOSSARY",
        "References",
    }
    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if text in section_starts:
            paragraph.paragraph_format.page_break_before = True


def remove_blank_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)
    paragraph._p = paragraph._element = None


def collapse_blanks_before_major_sections(doc):
    section_starts = {
        "Acknowledgement",
        "Abstract",
        "Chapter 1",
        "CHAPTER 2",
        "CHAPTER 3",
        "CHAPTER 4",
        "CHAPTER 5",
        "CHAPTER 6",
        "APPENDICES",
        "Appendix B: Machine Learning Model Configuration",
        "Appendix C: Sample Recommendation",
        "GLOSSARY",
        "References",
    }
    changed = True
    while changed:
        changed = False
        paragraphs = list(doc.paragraphs)
        for idx, paragraph in enumerate(paragraphs):
            text = " ".join(paragraph.text.split())
            if text not in section_starts:
                continue
            prior_idx = idx - 1
            while prior_idx >= 0:
                prior = paragraphs[prior_idx]
                prior_text = " ".join(prior.text.split())
                if prior_text or has_drawing(prior):
                    break
                remove_blank_paragraph(prior)
                changed = True
                prior_idx -= 1
            if changed:
                break


def normalize_table_captions(doc):
    captions_by_original = {
        "Table 3.1": TABLE_CAPTIONS[0],
        "Table 4.1": TABLE_CAPTIONS[1],
        "Table 5.1": TABLE_CAPTIONS[2],
        "Table 5.2": TABLE_CAPTIONS[3],
        "Table 5.3": TABLE_CAPTIONS[4],
        "Table 5.4": TABLE_CAPTIONS[5],
        "Table 5.5 Average Borrower Characteristics": TABLE_CAPTIONS[6],
        "Table 5.5 Overall": TABLE_CAPTIONS[7],
        "Table 5.6 Summary": TABLE_CAPTIONS[8],
        "Table A.1": TABLE_CAPTIONS[9],
        "Table A.2": TABLE_CAPTIONS[10],
        "Table A.3": TABLE_CAPTIONS[11],
        "Table B.1": TABLE_CAPTIONS[12],
        "Table B.2": TABLE_CAPTIONS[13],
    }
    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        for key, (number, title) in captions_by_original.items():
            if text.startswith(key):
                format_caption(paragraph, number, title)
                break

    table_index = 0
    blocks = list(iter_block_items(doc))
    for idx, block in enumerate(blocks):
        if not isinstance(block, Table):
            continue
        number, title = TABLE_CAPTIONS[table_index]
        table_index += 1
        prev = None
        for prior in reversed(blocks[:idx]):
            if isinstance(prior, Paragraph) and " ".join(prior.text.split()):
                prev = prior
                break
        prev_text = " ".join(prev.text.split()) if prev else ""
        if not prev_text.startswith("Table"):
            caption = paragraph_before_table(doc, block)
            format_caption(caption, number, title)
        block.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in block.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)


def normalize_figure_captions(doc):
    figure_index = 0
    blocks = list(iter_block_items(doc))
    chapter_three_idx = 0
    for idx, block in enumerate(blocks):
        if isinstance(block, Paragraph) and " ".join(block.text.split()) == "CHAPTER 3":
            chapter_three_idx = idx
            break
    for idx, block in enumerate(blocks):
        if not isinstance(block, Paragraph) or not has_drawing(block):
            continue
        if idx < chapter_three_idx:
            continue
        prev = None
        for prior in reversed(blocks[:idx]):
            if isinstance(prior, Paragraph) and " ".join(prior.text.split()):
                prev = prior
                break
        prev_text = " ".join(prev.text.split()) if prev else ""
        if figure_index >= len(FIGURE_CAPTIONS):
            continue
        number, title = FIGURE_CAPTIONS[figure_index]
        own_text = " ".join(block.text.split())
        if own_text.startswith("Figure"):
            clear_paragraph_text(block)
            caption = paragraph_after(doc, block)
            format_caption(caption, number, title)
            figure_index += 1
            continue
        if prev_text.startswith("Figure"):
            format_caption(prev, number, title)
            figure_index += 1
            continue
        caption = paragraph_after(doc, block)
        format_caption(caption, number, title)
        figure_index += 1


def add_list_table_after(doc, anchor, heading, rows):
    h = paragraph_after(doc, anchor, heading, style="Heading 1")
    h.paragraph_format.page_break_before = True
    h.paragraph_format.space_after = Pt(12)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.25)
    set_cell_text(table.rows[0].cells[0], "Number", bold=True)
    set_cell_text(table.rows[0].cells[1], "Title", bold=True)
    for number, title in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], number)
        set_cell_text(cells[1], title)
    h._p.addnext(table._tbl)
    return table


def add_front_matter_lists(doc):
    chapter_one = next(
        paragraph for paragraph in doc.paragraphs
        if " ".join(paragraph.text.split()) == "Chapter 1"
    )
    spacer = paragraph_before_paragraph(doc, chapter_one)
    table_listing = add_list_table_after(doc, spacer, "List of Tables", TABLE_CAPTIONS)
    list_of_figures_anchor = doc.add_paragraph()
    table_listing._tbl.addnext(list_of_figures_anchor._p)
    add_list_table_after(doc, list_of_figures_anchor, "List of Figures", FIGURE_CAPTIONS)


def main():
    doc = Document(SOURCE)
    ensure_caption_style(doc)
    normalize_styles(doc)
    normalize_table_captions(doc)
    normalize_figure_captions(doc)
    add_front_matter_lists(doc)
    collapse_blanks_before_major_sections(doc)
    set_major_section_page_breaks(doc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
